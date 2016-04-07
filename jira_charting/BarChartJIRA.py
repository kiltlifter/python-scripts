from datetime import date
import re
from openpyxl import Workbook
from openpyxl.chart import (
	BarChart,
	Series,
    Reference
)
from openpyxl.chart.axis import DateAxis

class Graph():
	def init(self):
		None

	def bar_chart(self, save_file_name, formatted_data, chart_title):
		wb = Workbook()
		ws = wb.active
	
		# Sample of what the format should be.
		"""
		rows = [
		    ['Date', 'Batch 1', 'Batch 2', 'Batch 3'],
		    [date(2015,9, 1), 40, 30, 25],
		    [date(2015,9, 2), 40, 25, 30],
		    [date(2015,9, 3), 50, 30, 45],
		    [date(2015,9, 4), 30, 25, 40],
		    [date(2015,9, 5), 25, 35, 30],
		    [date(2015,9, 6), 20, 40, 35],
		]
		"""
		rows = formatted_data
		for row in rows:
		    ws.append(tuple(row))
		
		title_regex = re.compile(r'\w+\s.*by\s(.*)')
		m = title_regex.match(chart_title)
		chart_type = m.groups()[0]

		chart1 = BarChart()
		chart1.type = "col"
		chart1.style = 10
		chart1.title = chart_title
		chart1.y_axis.title = 'Scale'
		chart1.x_axis.title = "%s Status" % chart_type

		chart1.height = 10.16
		chart1.width = 16.9418		
	
		data = Reference(ws, min_col=2, min_row=1, max_row=7, max_col=3)
		cats = Reference(ws, min_col=1, min_row=2, max_row=7)
		chart1.add_data(data, titles_from_data=True)
		chart1.set_categories(cats)
		chart1.shape = 4
		ws.add_chart(chart1, "E1")
		
		from copy import deepcopy

		chart2 = deepcopy(chart1)
		chart2.style = 11
		chart2.type = "bar"
		chart2.title = chart_title
		chart2.height = 10.16
		chart2.width = 16.9418	

		ws.add_chart(chart2, "E21")

		wb.save(save_file_name + ".xlsx")
