
from openpyxl import Workbook

from openpyxl.chart import (
    PieChart,
    Reference
)
from openpyxl.chart.series import DataPoint

class PieGraph():
	def __init__(self):
		None


	def chart(self, save_file_path, chart_data, graph_title, point_total=None):
		data = chart_data
		
		wb = Workbook()
		ws = wb.active
		
		for row in data:
			ws.append(row)
		
		pie = PieChart()
		labels = Reference(ws, min_col=1, min_row=2, max_row=len(data))
		data = Reference(ws, min_col=2, min_row=1, max_row=len(data))
		pie.add_data(data, titles_from_data=True)
		pie.set_categories(labels)
		pie.title = graph_title
		pie.dataLabels
		pie.height = 10.16
		pie.width = 16.9418		
		# Cut the first slice out of the pie
		#slice = DataPoint(idx=0, explosion=20)
		#pie.series[0].data_points = [slice]
		
		ws.add_chart(pie, "D1")
		ws["C1"] = point_total
		wb.save(save_file_path + ".xlsx")
